from flask import Flask, render_template, request, jsonify, session
from flask_socketio import SocketIO, emit, join_room
import json, os, uuid, threading, time, random, hashlib

app = Flask(__name__)
app.config['SECRET_KEY'] = 'ipl_auction_secret_2024'
socketio = SocketIO(app, cors_allowed_origins="*", async_mode='threading')

# ── State ─────────────────────────────────────────────────────────────────────
rooms       = {}
sid_to_room = {}
color_pool  = ['#e63946','#2a9d8f','#e9c46a','#f4a261','#264653',
               '#8338ec','#3a86ff','#fb5607','#06d6a0','#118ab2']

INITIAL_BUDGET      = 10000
TIMER_SECONDS       = 20
MIN_TIMER_AFTER_BID = 10

SET_LABELS = {
    'Batsman':       'SET A — BATSMEN',
    'Wicket-Keeper': 'SET B — WICKET-KEEPERS',
    'All-Rounder':   'SET C — ALL-ROUNDERS',
    'Bowler':        'SET D — BOWLERS',
}

DATA_DIR   = os.path.join(os.path.dirname(__file__), 'data')
USERS_FILE = os.path.join(DATA_DIR, 'users.json')

# ── User helpers ──────────────────────────────────────────────────────────────
def load_users():
    try:
        with open(USERS_FILE) as f: return json.load(f)
    except: return {}

def save_users(users):
    with open(USERS_FILE, 'w') as f: json.dump(users, f, indent=2)

def hash_pw(pw):
    return hashlib.sha256(pw.encode()).hexdigest()

# ── Player helpers ─────────────────────────────────────────────────────────────
def load_players():
    xlsx = os.path.join(DATA_DIR, 'ipl_players_2025.xlsx')
    if os.path.exists(xlsx):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(xlsx, data_only=True)
            ws = wb['Players']
            headers = [str(ws.cell(3,c).value).strip().lower()
                       .replace(' ','_').replace('(₹l)','').strip()
                       for c in range(1, ws.max_column+1)]
            players = []
            for row in ws.iter_rows(min_row=4, values_only=True):
                if not row[1]: continue
                p = dict(zip(headers, row))
                players.append({
                    'id':          p.get('#') or len(players)+1,
                    'name':        str(p.get('name','')).strip(),
                    'role':        str(p.get('role','')).strip(),
                    'nationality': str(p.get('nationality','India')).strip(),
                    'team':        str(p.get('ipl_franchise') or p.get('franchise','—')).strip(),
                    'base_price':  int(p.get('base_price') or p.get('base_price_') or 75),
                    'rating':      int(p.get('rating') or 70),
                })
            return players
        except Exception as e:
            print(f'Excel load error: {e}')
    with open(os.path.join(DATA_DIR, 'players.json')) as f:
        return json.load(f)

def arrange_in_sets(players):
    order = ['Batsman','Wicket-Keeper','All-Rounder','Bowler']
    buckets = {r: [] for r in order}
    for p in players:
        buckets.setdefault(p.get('role','Batsman'), []).append(p)
    result = []
    for role in order:
        group = buckets.get(role, [])
        random.shuffle(group)
        for i, p in enumerate(group):
            p['_set_start'] = (i == 0)
            p['_set_label'] = SET_LABELS.get(role, role)
        result.extend(group)
    return result

# ── Room helpers ───────────────────────────────────────────────────────────────
def new_room(room_id, host_sid, host_name):
    raw = load_players()
    players = arrange_in_sets(raw)
    return {
        'id': room_id, 'host_sid': host_sid, 'host_name': host_name,
        'status': 'waiting', 'teams': {}, 'players': players,
        'player_index': -1, 'current_player': None,
        'current_bid': 0, 'current_bidder_sid': None, 'current_bidder_name': None,
        'timer': TIMER_SECONDS, 'timer_lock': threading.Lock(),
        'timer_stop': threading.Event(), 'sold_players': [], 'unsold_players': [],
        'started_at': time.time(),
    }

def new_team(sid, name, idx):
    return {'sid': sid, 'name': name, 'budget': INITIAL_BUDGET,
            'players': [], 'color': color_pool[idx % len(color_pool)]}

def pub(room):
    idx = room['player_index']
    return {
        'id': room['id'], 'host_sid': room['host_sid'], 'host_name': room['host_name'],
        'status': room['status'], 'teams': list(room['teams'].values()),
        'current_player': room['current_player'], 'current_bid': room['current_bid'],
        'current_bidder_sid': room['current_bidder_sid'],
        'current_bidder_name': room['current_bidder_name'],
        'timer': room['timer'],
        'players_left': max(0, len(room['players']) - idx - 1),
        'total_players': len(room['players']),
        'sold_count': len(room['sold_players']),
        'upcoming_players': room['players'][idx+1 : idx+8],
    }

# ── Timer ──────────────────────────────────────────────────────────────────────
def _timer_worker(room_id, stop_event):
    while not stop_event.is_set():
        time.sleep(1)
        if stop_event.is_set(): break
        room = rooms.get(room_id)
        if not room or room['status'] != 'bidding': break
        with room['timer_lock']:
            if room['status'] != 'bidding': break
            room['timer'] -= 1
            t = room['timer']
        socketio.emit('timer_tick', {'timer': t}, to=room_id)
        if t <= 0:
            _close_bid(room_id); break

def _start_timer(room):
    room['timer_stop'].set()
    stop = threading.Event()
    room['timer_stop'] = stop
    room['timer'] = TIMER_SECONDS
    threading.Thread(target=_timer_worker, args=(room['id'], stop), daemon=True).start()

def _close_bid(room_id):
    room = rooms.get(room_id)
    if not room or room['status'] != 'bidding': return
    player = room['current_player']
    if room['current_bidder_sid']:
        winner = room['teams'].get(room['current_bidder_sid'])
        if winner:
            winner['budget'] -= room['current_bid']
            winner['players'].append({**player, 'sold_price': room['current_bid']})
        room['sold_players'].append({
            **player, 'sold_price': room['current_bid'],
            'sold_to': room['current_bidder_name'],
            'sold_to_color': room['teams'].get(room['current_bidder_sid'],{}).get('color','#fff'),
        })
        room['status'] = 'sold'
        socketio.emit('player_sold', {
            'player': player, 'price': room['current_bid'],
            'team': room['current_bidder_name'],
            'color': room['teams'].get(room['current_bidder_sid'],{}).get('color','#fff'),
            'room': pub(room),
        }, to=room_id)
    else:
        room['unsold_players'].append(player)
        room['status'] = 'unsold'
        socketio.emit('player_unsold', {'player': player, 'room': pub(room)}, to=room_id)

def _finish_auction(room_id):
    room = rooms.get(room_id)
    if not room: return
    room['status'] = 'finished'
    room['timer_stop'].set()
    # Save game record for logged-in users
    _save_game_records(room)
    socketio.emit('auction_finished', {'room': pub(room), 'sold': room['sold_players']}, to=room_id)

def _save_game_records(room):
    """Persist auction results for any logged-in participants."""
    users = load_users()
    changed = False
    for sid, team in room['teams'].items():
        username = team.get('username')
        if not username or username not in users: continue
        record = {
            'room_id':    room['id'],
            'date':       time.strftime('%Y-%m-%d %H:%M'),
            'team_name':  team['name'],
            'budget_left':team['budget'],
            'players_won':len(team['players']),
            'total_spend':INITIAL_BUDGET - team['budget'],
            'squad':      [{'name':p['name'],'role':p['role'],'price':p['sold_price']} for p in team['players']],
            'players_total': len(room['players']),
            'sold_count': len(room['sold_players']),
        }
        users[username].setdefault('records', []).insert(0, record)
        users[username]['records'] = users[username]['records'][:20]  # keep last 20
        changed = True
    if changed: save_users(users)

def _next_player(room_id):
    room = rooms.get(room_id)
    if not room: return
    room['player_index'] += 1
    if room['player_index'] >= len(room['players']):
        _finish_auction(room_id); return
    player = room['players'][room['player_index']]
    room['current_player'] = player
    room['current_bid'] = player['base_price']
    room['current_bidder_sid'] = None
    room['current_bidder_name'] = None
    room['status'] = 'bidding'
    _start_timer(room)
    socketio.emit('new_player', {'player': player, 'room': pub(room)}, to=room_id)

# ── HTTP routes ────────────────────────────────────────────────────────────────
@app.route('/')
def index():
    return render_template('index.html')

@app.route('/api/rooms')
def api_rooms():
    return jsonify([
        {'id': r['id'], 'host': r['host_name'], 'teams': len(r['teams']), 'status': r['status']}
        for r in rooms.values() if r['status'] != 'finished'
    ])

# ── Auth API ───────────────────────────────────────────────────────────────────
@app.route('/api/auth/register', methods=['POST'])
def register():
    data = request.get_json() or {}
    username = (data.get('username') or '').strip().lower()
    password = (data.get('password') or '').strip()
    display  = (data.get('display') or username).strip()[:20]
    if not username or not password:
        return jsonify({'ok': False, 'msg': 'Username and password required'}), 400
    if len(username) < 3:
        return jsonify({'ok': False, 'msg': 'Username must be at least 3 characters'}), 400
    users = load_users()
    if username in users:
        return jsonify({'ok': False, 'msg': 'Username already taken'}), 409
    users[username] = {
        'username': username, 'display': display,
        'password': hash_pw(password),
        'created':  time.strftime('%Y-%m-%d'),
        'records':  [],
    }
    save_users(users)
    return jsonify({'ok': True, 'user': {'username': username, 'display': display}})

@app.route('/api/auth/login', methods=['POST'])
def login():
    data = request.get_json() or {}
    username = (data.get('username') or '').strip().lower()
    password = (data.get('password') or '').strip()
    users = load_users()
    u = users.get(username)
    if not u or u['password'] != hash_pw(password):
        return jsonify({'ok': False, 'msg': 'Invalid username or password'}), 401
    return jsonify({'ok': True, 'user': {'username': username, 'display': u.get('display', username)}})

@app.route('/api/profile/<username>')
def get_profile(username):
    users = load_users()
    u = users.get(username.lower())
    if not u: return jsonify({'ok': False, 'msg': 'User not found'}), 404
    return jsonify({'ok': True, 'display': u.get('display', username),
                    'created': u.get('created',''), 'records': u.get('records',[])})

# ── Socket events ──────────────────────────────────────────────────────────────
@socketio.on('connect')
def on_connect():
    emit('connected', {'sid': request.sid})

@socketio.on('disconnect')
def on_disconnect():
    sid = request.sid
    room_id = sid_to_room.pop(sid, None)
    if not room_id: return
    room = rooms.get(room_id)
    if not room: return
    leaving = room['teams'].pop(sid, {})
    name = leaving.get('name', '?')
    if room['current_bidder_sid'] == sid:
        room['current_bidder_sid'] = None
        room['current_bidder_name'] = None
    if not room['teams']:
        room['timer_stop'].set(); rooms.pop(room_id, None); return
    socketio.emit('notification', {'msg': f'👋 {name} left'}, to=room_id)
    socketio.emit('room_update', pub(room), to=room_id)

@socketio.on('create_room')
def on_create_room(data):
    sid  = request.sid
    name = (data.get('name') or 'Host').strip()[:20]
    username = data.get('username','')
    room_id = str(uuid.uuid4())[:6].upper()
    room = new_room(room_id, sid, name)
    team = new_team(sid, name, 0)
    if username: team['username'] = username
    room['teams'][sid] = team
    rooms[room_id]   = room
    sid_to_room[sid] = room_id
    join_room(room_id)
    emit('room_joined', {'room': pub(room), 'your_sid': sid, 'room_id': room_id})

@socketio.on('join_room_req')
def on_join_room(data):
    sid     = request.sid
    room_id = (data.get('room_id') or '').strip().upper()
    name    = (data.get('name') or 'Team').strip()[:20]
    username = data.get('username','')
    room = rooms.get(room_id)
    if not room:
        emit('error', {'msg': f'Room "{room_id}" not found.'}); return
    if room['status'] == 'finished':
        emit('error', {'msg': 'This auction has ended.'}); return
    idx  = len(room['teams'])
    team = new_team(sid, name, idx)
    if username: team['username'] = username
    room['teams'][sid] = team
    sid_to_room[sid] = room_id
    join_room(room_id)
    emit('room_joined', {'room': pub(room), 'your_sid': sid, 'room_id': room_id})
    socketio.emit('room_update', pub(room), to=room_id)
    socketio.emit('notification', {'msg': f'🏏 {name} joined!'}, to=room_id)

@socketio.on('start_auction')
def on_start(data):
    sid = request.sid; room_id = sid_to_room.get(sid); room = rooms.get(room_id)
    if not room: emit('error',{'msg':'Not in a room.'}); return
    if room['host_sid'] != sid: emit('error',{'msg':'Only the host can start.'}); return
    if room['status'] != 'waiting': emit('error',{'msg':'Auction already started.'}); return
    _next_player(room_id)

@socketio.on('next_player')
def on_next(data):
    sid = request.sid; room_id = sid_to_room.get(sid); room = rooms.get(room_id)
    if not room or room['host_sid'] != sid: return
    if room['status'] not in ('sold','unsold','waiting'):
        emit('error',{'msg':'Cannot skip right now.'}); return
    _next_player(room_id)

@socketio.on('stop_auction')
def on_stop(data):
    sid = request.sid; room_id = sid_to_room.get(sid); room = rooms.get(room_id)
    if not room or room['host_sid'] != sid: return
    if room['status'] == 'finished': return
    room['timer_stop'].set()
    room['status'] = 'sold'
    _finish_auction(room_id)

@socketio.on('place_bid')
def on_bid(data):
    sid = request.sid; room_id = sid_to_room.get(sid); room = rooms.get(room_id)
    if not room or room['status'] != 'bidding': return
    team = room['teams'].get(sid)
    if not team: return
    try: amount = int(data.get('amount', 0))
    except: return
    if amount <= 0: return
    new_bid = room['current_bid'] + amount
    if new_bid > team['budget']:
        emit('error',{'msg':f'Not enough budget! You have ₹{team["budget"]}L left.'}); return
    if sid == room['current_bidder_sid']:
        emit('error',{'msg':"You're already highest bidder!"}); return
    with room['timer_lock']:
        room['current_bid'] = new_bid
        room['current_bidder_sid'] = sid
        room['current_bidder_name'] = team['name']
        room['timer'] = max(room['timer'], MIN_TIMER_AFTER_BID)
    socketio.emit('bid_update', {
        'bid': new_bid, 'bidder': team['name'],
        'bidder_color': team['color'], 'bidder_sid': sid,
        'timer': room['timer'], 'room': pub(room),
    }, to=room_id)

@socketio.on('get_state')
def on_get_state(data):
    sid = request.sid; room_id = sid_to_room.get(sid); room = rooms.get(room_id)
    if room: emit('full_state', pub(room))

if __name__ == '__main__':
    socketio.run(app, host='0.0.0.0', port=5000, debug=False, use_reloader=False)
